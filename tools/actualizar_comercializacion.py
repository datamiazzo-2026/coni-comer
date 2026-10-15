#!/usr/bin/env python3
"""
Actualiza el bloque EMBEDDED de dashboard/index.html a partir de 'Comercialización.xlsx'.
Uso:  python3 tools/actualizar_comercializacion.py <ruta.xlsx> [dashboard/index.html] [--dry]
- Lee el Excel de forma estructurada (openpyxl), sin regex sobre texto.
- Compras-Ventas: para soja, maíz y trigo toma, de cada (campaña, mes), la ÚLTIMA fila del archivo
  (el último bloque es el alineado a la semana vigente) y actualiza:
    * todas las filas del último bloque (campaña vigente + semana alineada de campañas anteriores)
    * todas las filas de la campaña vigente.
- Liquidación: reutiliza el parser JS del propio tablero (parseLiquidacion) sobre la hoja del Excel.
- Idempotente: si no hay novedades no modifica nada.
Imprime un reporte de cambios. Código de salida 0 = ok.
"""
import sys, re, json, subprocess, datetime, tempfile, os
import openpyxl

args = [a for a in sys.argv[1:] if not a.startswith('--')]
DRY = '--dry' in sys.argv
xlsx = args[0]; html_path = args[1] if len(args) > 1 else 'dashboard/index.html'

MES = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
LAYOUT = {  # columna (1-based) de producción y si lleva "queda + por fijar"
    'soja':  {'off': 2,  'extra': True},
    'maiz':  {'off': 17, 'extra': False},
    'trigo': {'off': 31, 'extra': False},
}

def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None
def r1(v):
    v = num(v)
    if v is None: return None
    x = round(v + 1e-9, 1)
    return int(x) if x == int(x) else x
def r0(v):
    v = num(v)
    return None if v is None else int(v + 0.5) if v >= 0 else -int(-v + 0.5)
def pct(v):
    v = num(v)
    return None if v is None else int(v * 100 + 0.5)

wb = openpyxl.load_workbook(xlsx, data_only=True)
ws = wb['Compras - Ventas']

# --- bloques: filas de etiqueta (mes) en la columna A ---
label_rows = [r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, str) and ws.cell(r, 1).value.strip().lower() in
              ('enero','febrero','marzo','abril','mayo','junio','julio','agosto','septiembre','octubre','noviembre','diciembre')]
last_label = label_rows[-1]

def record(r, crop):
    L = LAYOUT[crop]; o = L['off']
    g = lambda k: ws.cell(r, o + k).value
    prod, mes, fecha, camp = num(g(0)), g(1), g(2), g(3)
    if prod is None or not isinstance(mes, datetime.datetime) or not isinstance(fecha, datetime.datetime) or not isinstance(camp, str) or not re.match(r'^\d\d/\d\d$', camp.strip()):
        return None
    rec = {'campaña': camp.strip(), 'mes_ano': f'{MES[mes.month-1]}-{str(mes.year)[2:]}',
           'fecha_info': f'{fecha.month}/{fecha.day}/{fecha.year}',
           'produccion': r1(prod), 'pct_comprado': pct(g(4)), 'total_comprado': r1(g(5)), 'expo': r1(g(6)), 'industria': r1(g(7))}
    if L['extra']:
        rec.update({'a_fijar_1': r1(g(8)), 'a_fijar_2': pct(g(9)), 'queda_vender_tn': r0(g(10)), 'queda_vender_mas_fijar': r0(g(11))})
    else:
        rec.update({'saldo_fijar_1': r1(g(8)), 'saldo_fijar_2': pct(g(9)), 'queda_vender_tn': r0(g(10))})
    rec['precio'] = None; rec['queda_vender_musd'] = None
    if rec['total_comprado'] is None or rec['queda_vender_tn'] is None:
        return None
    return rec

latest = {c: {} for c in LAYOUT}      # (campaña, mes_ano) -> (fila, registro) última aparición
for c in LAYOUT:
    for r in range(8, ws.max_row + 1):
        rec = record(r, c)
        if rec: latest[c][(rec['campaña'], rec['mes_ano'])] = (r, rec)

vigente = max(k[0] for k in latest['soja'])
html = open(html_path, encoding='utf-8').read()
m = re.search(r'const EMBEDDED = (\{.*?\});\n', html, re.S)
E = json.loads(m.group(1))

report = []
for c in LAYOUT:
    arr = E['raw'][c]
    for key, (row, rec) in latest[c].items():
        if not (row >= last_label or key[0] == vigente):
            continue
        idx = [i for i, x in enumerate(arr) if (x['campaña'], x['mes_ano']) == key]
        if idx:
            old = arr[idx[0]]
            newrec = {k: rec.get(k) for k in old.keys()}
            for k in rec:                       # claves nuevas: conservar
                if k not in newrec: newrec[k] = rec[k]
            newrec['precio'] = old.get('precio'); newrec['queda_vender_musd'] = old.get('queda_vender_musd')
            def distinto(a, b):   # la planilla muestra a veces enteros y a veces 1 decimal: se toleran diferencias de redondeo de visualización (<=0.5)
                if isinstance(a, (int, float)) and isinstance(b, (int, float)): return abs(a - b) > 0.5
                return a != b
            for k in newrec:
                if not distinto(old.get(k), newrec.get(k)): newrec[k] = old.get(k)
            if any(old.get(k) != newrec.get(k) for k in newrec):
                diff = {k: (old.get(k), newrec.get(k)) for k in newrec if old.get(k) != newrec.get(k)}
                report.append(f'CAMBIA  {c} {key[0]} {key[1]}: {diff}')
                for i in idx: arr[i] = newrec
        else:
            arr.append(rec); report.append(f'NUEVA   {c} {key[0]} {key[1]} fecha {rec["fecha_info"]} queda {rec["queda_vender_tn"]}')

# --- Liquidación: reutiliza parseLiquidacion del tablero ---
wl = wb['Liquidación']
def fmt(v):
    if v is None: return ''
    if isinstance(v, (int, float)):
        if float(v) == int(v): return str(int(v))
        return '"' + repr(float(v)).replace('.', ',') + '"'
    return str(v).replace('"', '""').replace(',', ' ')
lines = []
for r in range(1, wl.max_row + 1):
    lines.append(','.join(fmt(wl.cell(r, c).value) for c in range(1, 45)))
csv_text = '\n'.join(lines)
a = html.index('function parseCSV'); b = html.index('async function loadLiveData')
js = html[a:b] + "\nconst fs=require('fs');const liq=parseLiquidacion(fs.readFileSync(process.argv[2],'utf8'));console.log(JSON.stringify(liq));"
with tempfile.TemporaryDirectory() as td:
    open(os.path.join(td, 'liq.csv'), 'w', encoding='utf-8').write(csv_text)
    open(os.path.join(td, 'p.js'), 'w', encoding='utf-8').write(js)
    out = subprocess.run(['node', os.path.join(td, 'p.js'), os.path.join(td, 'liq.csv')], capture_output=True, text=True)
if out.returncode != 0:
    print('ERROR parseando liquidación:', out.stderr[:500]); sys.exit(2)
liq = json.loads(out.stdout)
old_liq = E['liq']
def liq_diff(a, b):
    out = []
    def igual(x, y):
        if isinstance(x, dict) and isinstance(y, dict):
            return set(x) == set(y) and all(igual(x[k], y[k]) for k in x)
        if isinstance(x, list) and isinstance(y, list):
            return len(x) == len(y) and all(igual(i, j) for i, j in zip(x, y))
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            return abs(x - y) <= 0.5
        return x == y
    for k in set(a) | set(b):
        if not igual(a.get(k), b.get(k)): out.append(k)
    return out
dl = liq_diff(liq, old_liq)
if liq['cutoffMonthIdx'] >= old_liq['cutoffMonthIdx'] and dl:
    report.append(f'(campos distintos en liquidación: {dl})')
    report.append(f'LIQUIDACIÓN actualizada: hasta {old_liq["cutoffMonthName"]} -> {liq["cutoffMonthName"]}; 2026 = {liq["mensual"].get("2026")}')
    E['liq'] = liq
elif liq['cutoffMonthIdx'] < old_liq['cutoffMonthIdx']:
    report.append(f'AVISO: el Excel trae liquidación hasta {liq["cutoffMonthName"]}, MÁS ATRÁS que la embebida ({old_liq["cutoffMonthName"]}). No se modifica.')

print(f'Campaña vigente: {vigente} | último bloque del Excel desde fila {last_label}')
print('\n'.join(report) if report else 'SIN CAMBIOS: el tablero ya tiene los datos del Excel.')
if report and not DRY:
    html = html[:m.start(1)] + json.dumps(E, separators=(',', ':'), ensure_ascii=False) + html[m.end(1):]
    open(html_path, 'w', encoding='utf-8').write(html)
    print('index.html actualizado.')
